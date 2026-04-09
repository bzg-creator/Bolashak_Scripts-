
import pyodbc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


DB_SERVER   = ""
DB_DATABASE = ""


NAME_MAP = {
    "США":                                    "Соединённые Штаты Америки",
    "USA":                                    "Соединённые Штаты Америки",
    "Америка":                                "Соединённые Штаты Америки",
    "United States":                          "Соединённые Штаты Америки",
    "United States of America":              "Соединённые Штаты Америки",
    "Соединенные Штаты Америки":             "Соединённые Штаты Америки",
    "Америка Құрама Штаттары":               "Соединённые Штаты Америки",
    "Америка құрама штаттары":               "Соединённые Штаты Америки",
    "Россия":                                "Российская Федерация",
    "РФ":                                    "Российская Федерация",
    "Ресей":                                 "Российская Федерация",
    "Ресей Федерациясы":                     "Российская Федерация",
    "Казахстан":                             "Республика Казахстан",
    "Қазақстан":                             "Республика Казахстан",
    "Германия":                              "Федеративная Республика Германия",
    "Germany":                               "Федеративная Республика Германия",
    "Deutschland":                           "Федеративная Республика Германия",
    "Германия Федеративті Республикасы":     "Федеративная Республика Германия",
    "Федеративная Республика Германии":      "Федеративная Республика Германия",
    "Великобритания":                        "Соединённое Королевство Великобритании и Северной Ирландии",
    "Соединенное Королевство Великобритании и Северной Ирландии":
                                             "Соединённое Королевство Великобритании и Северной Ирландии",
    "Ұлыбритания және Солтүстік Ирландия Біріккен Патшалығы":
                                             "Соединённое Королевство Великобритании и Северной Ирландии",
    "500 Великобритания":                    "Соединённое Королевство Великобритании и Северной Ирландии",
    "Узбекистан":                            "Республика Узбекистан",
    "Белоруссия":                            "Республика Беларусь",
    "Беларусь":                              "Республика Беларусь",
    "Беларусь Республикасы":                 "Республика Беларусь",
    "Южная Корея":                           "Республика Корея",
    "Корея":                                 "Республика Корея",
    "Турция":                                "Турецкая Республика",
    "Түркия":                                "Турецкая Республика",
    "Нидерланды":                            "Королевство Нидерланды",
    "Голландия":                             "Королевство Нидерланды",
    "Швеция":                                "Королевство Швеция",
    "Швеция Патшалығы":                      "Королевство Швеция",
    "Норвегия":                              "Королевство Норвегия",
    "Дания":                                 "Королевство Дания",
    "Испания":                               "Королевство Испания",
    "Бельгия":                               "Королевство Бельгия",
    "Польша":                                "Республика Польша",
    "Чехия":                                 "Чешская Республика",
    "Италия":                                "Итальянская Республика",
    "Венгрия":                               "Республика Венгрия",
    "Азербайджан":                           "Азербайджанская Республика",
    "Армения":                               "Республика Армения",
    "Государство Израиль":                   "Израиль",
    "Ирландия":                              "Ирландская Республика",
    "Эстония":                               "Эстонская Республика",
    "Швейцария":                             "Швейцарская Конфедерация",
    "Франция":                               "Французская Республика",
}

# ─────────────────────────────────────────────
# Excel стили
# ─────────────────────────────────────────────
HDR_FONT    = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HDR_FILL_1  = PatternFill("solid", fgColor="1F4E79")   # синий  — уровень 1
HDR_FILL_2  = PatternFill("solid", fgColor="375623")   # зелёный — уровень 2
HDR_FILL_3  = PatternFill("solid", fgColor="7B3F00")   # коричневый — уровень 3
HDR_FILL_S  = PatternFill("solid", fgColor="2E4057")   # тёмный — сводка
HDR_ALIGN   = Alignment(horizontal="center", vertical="center", wrap_text=True)
NORM_FONT   = Font(name="Arial", size=10)
BOLD_FONT   = Font(name="Arial", size=10, bold=True)
WARN_FILL   = PatternFill("solid", fgColor="FFF2CC")
ERR_FILL    = PatternFill("solid", fgColor="FCE4D6")
OK_FILL     = PatternFill("solid", fgColor="E2EFDA")
BLUE_FILL   = PatternFill("solid", fgColor="DEEAF1")
thin        = Side(style="thin", color="D9D9D9")
BORDER      = Border(left=thin, right=thin, top=thin, bottom=thin)


def get_conn():
    return pyodbc.connect(
        "DRIVER={ODBC Driver 17 for SQL Server};"
        f"SERVER={DB_SERVER};DATABASE={DB_DATABASE};"
        "Trusted_Connection=yes;",
        timeout=0
    )


def style_header(ws, row, ncols, fill):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font    = HDR_FONT
        c.fill    = fill
        c.alignment = HDR_ALIGN
        c.border  = BORDER
    ws.row_dimensions[row].height = 30


def write_cell(ws, row, col, value, fill=None, bold=False, center=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font   = BOLD_FONT if bold else NORM_FONT
    c.border = BORDER
    c.alignment = Alignment(
        vertical="center",
        horizontal="center" if center else "left",
        wrap_text=True
    )
    if fill:
        c.fill = fill
    return c


def add_title(ws, text, ncols, color="1F4E79"):
    ws.merge_cells(f"A1:{get_column_letter(ncols)}1")
    c = ws["A1"]
    c.value     = text
    c.font      = Font(name="Arial", bold=True, size=12, color=color)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill      = PatternFill("solid", fgColor="F2F2F2")
    ws.row_dimensions[1].height = 22


# ============================================================
# ЗАПРОСЫ К БД
# ============================================================

def fetch_level1(cursor):
    """Уровень 1: названия в Dictionaries которые нужно исправить."""
    print("  Уровень 1: получаем названия из справочника...", end=" ")

    cursor.execute("""
        SELECT
            CONVERT(NVARCHAR(36), d.Id) AS Id,
            d.[Name]    AS CurrentName,
            d.Code,
            d.IsDeleted,
            d.IsCurrent,
            d.[Type]
        FROM dbo.Dictionaries d
        WHERE d.[Type] = N'CountryDictionary'
          AND d.IsDeleted = 0
        ORDER BY d.[Name]
    """)
    rows = cursor.fetchall()
    print(f"{len(rows)} строк")

    result = []
    for row in rows:
        guid, name, code, is_del, is_cur, dtype = row
        name_str = (name or "").strip()
        correct  = NAME_MAP.get(name_str)
        if correct and correct != name_str:
            result.append({
                "guid":     guid,
                "current":  name_str,
                "correct":  correct,
                "code":     code,
                "is_del":   is_del,
                "is_cur":   is_cur,
            })
    return result


def fetch_level2(cursor):
    """Уровень 2: GUID дубли — одна страна, несколько GUID."""
    print("  Уровень 2: ищем дублирующиеся GUID...", end=" ")

    # Находим все GUID которые ссылаются на одно и то же правильное название
    cursor.execute("""
        SELECT
            CONVERT(NVARCHAR(36), d.Id) AS Id,
            LTRIM(RTRIM(d.[Name]))       AS Name,
            d.IsCurrent,
            d.IsDeleted,
            d.Code
        FROM dbo.Dictionaries d
        WHERE d.[Type] = N'CountryDictionary'
        ORDER BY d.[Name], d.IsCurrent DESC
    """)
    all_dicts = cursor.fetchall()
    print(f"{len(all_dicts)} записей в справочнике")

    # Нормализуем имена и группируем
    from collections import defaultdict
    groups = defaultdict(list)
    for row in all_dicts:
        guid, name, is_cur, is_del, code = row
        norm = NAME_MAP.get(name, name)
        groups[norm].append({
            "guid": guid, "name": name,
            "is_cur": is_cur, "is_del": is_del, "code": code
        })

    # Оставляем только группы с дублями
    dupes = {k: v for k, v in groups.items() if len(v) > 1}

    # Для каждого дубля считаем записи быстро (по Id)
    result = []
    tables = [
        ("Cards.Country",         "SELECT COUNT(*) FROM dbo.Cards WITH(NOLOCK) WHERE CountryDictionaryId=?"),
        ("Cards.Commission",      "SELECT COUNT(*) FROM dbo.Cards WITH(NOLOCK) WHERE CommissionCountryDictionaryId=?"),
        ("Cards.LangSchool",      "SELECT COUNT(*) FROM dbo.Cards WITH(NOLOCK) WHERE LanguageSchoolCountryDictionaryId=?"),
        ("Placements",            "SELECT COUNT(*) FROM dbo.Placements WITH(NOLOCK) WHERE CountryDictionaryId=?"),
        ("Dogovors",              "SELECT COUNT(*) FROM dbo.Dogovors WITH(NOLOCK) WHERE CountryDictionaryId=?"),
        ("Movements",             "SELECT COUNT(*) FROM dbo.Movements WITH(NOLOCK) WHERE CountryDictionaryId=?"),
        ("Plans",                 "SELECT COUNT(*) FROM dbo.Plans WITH(NOLOCK) WHERE CountryDictionaryId=?"),
    ]

    for correct_name, entries in dupes.items():
        # Определяем правильный GUID (IsCurrent=1, IsDeleted=0)
        correct_entry = next(
            (e for e in entries if e["is_cur"] and not e["is_del"]), entries[0]
        )
        wrong_entries = [e for e in entries if e["guid"] != correct_entry["guid"]]

        for wrong in wrong_entries:
            counts = {}
            for tname, sql in tables:
                cursor.execute(sql, wrong["guid"])
                counts[tname] = cursor.fetchone()[0]
            total = sum(counts.values())
            result.append({
                "correct_name":  correct_name,
                "wrong_guid":    wrong["guid"],
                "wrong_name":    wrong["name"],
                "correct_guid":  correct_entry["guid"],
                "correct_name2": correct_entry["name"],
                "total":         total,
                "counts":        counts,
            })

    return result


def fetch_level3(cursor):
    """Уровень 3: записи без Id (NULL) которые нужно привязать к GUID."""
    print("  Уровень 3: ищем записи без Id...", end=" ")

    # Сначала получаем справочник для матчинга
    cursor.execute("""
        SELECT CONVERT(NVARCHAR(36), d.Id), LTRIM(RTRIM(d.[Name]))
        FROM dbo.Dictionaries d
        WHERE d.[Type] = N'CountryDictionary'
          AND d.IsCurrent = 1 AND d.IsDeleted = 0
    """)
    dict_lookup = {name: guid for guid, name in cursor.fetchall()}

    # Запрос по всем таблицам кроме Movements (там нет смысла — только Id)
    sql = """
        SELECT src.[Value], src.Таблица, COUNT(*) AS Кол
        FROM (
            SELECT NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N'')             AS [Value], N'Cards.Country'       AS Таблица FROM dbo.Cards      WHERE NULLIF(LTRIM(RTRIM(CountryDictionaryId)),N'') IS NULL AND NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N'') IS NOT NULL
            UNION ALL
            SELECT NULLIF(LTRIM(RTRIM(CommissionCountryDictionaryValue)),N''),              N'Cards.Commission'    FROM dbo.Cards      WHERE NULLIF(LTRIM(RTRIM(CommissionCountryDictionaryId)),N'') IS NULL AND NULLIF(LTRIM(RTRIM(CommissionCountryDictionaryValue)),N'') IS NOT NULL
            UNION ALL
            SELECT NULLIF(LTRIM(RTRIM(LanguageSchoolCountryDictionaryValue)),N''),          N'Cards.LangSchool'    FROM dbo.Cards      WHERE NULLIF(LTRIM(RTRIM(LanguageSchoolCountryDictionaryId)),N'') IS NULL AND NULLIF(LTRIM(RTRIM(LanguageSchoolCountryDictionaryValue)),N'') IS NOT NULL
            UNION ALL
            SELECT NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N''),                        N'Placements'          FROM dbo.Placements WHERE NULLIF(LTRIM(RTRIM(CountryDictionaryId)),N'') IS NULL AND NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N'') IS NOT NULL
            UNION ALL
            SELECT NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N''),                        N'Dogovors'            FROM dbo.Dogovors   WHERE NULLIF(LTRIM(RTRIM(CountryDictionaryId)),N'') IS NULL AND NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N'') IS NOT NULL
            UNION ALL
            SELECT NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N''),                        N'Plans'               FROM dbo.Plans      WHERE NULLIF(LTRIM(RTRIM(CountryDictionaryId)),N'') IS NULL AND NULLIF(LTRIM(RTRIM(CountryDictionaryValue)),N'') IS NOT NULL
        ) src
        GROUP BY src.[Value], src.Таблица
        ORDER BY Кол DESC
    """
    cursor.execute(sql)
    rows = cursor.fetchall()
    print(f"{len(rows)} строк")

    result = []
    for value, table, cnt in rows:
        norm    = NAME_MAP.get((value or "").strip(), (value or "").strip())
        guid    = dict_lookup.get(norm)
        status  = "Найден GUID" if guid else "GUID не найден"
        result.append({
            "value":   value,
            "table":   table,
            "count":   cnt,
            "correct": norm,
            "guid":    guid or "—",
            "status":  status,
        })
    return result


# ============================================================
# ПОСТРОЕНИЕ EXCEL
# ============================================================

def build_excel(l1, l2, l3, output_path):
    wb = openpyxl.Workbook()

    # ── Лист 1: Уровень 1 — исправить названия ───────────────
    ws1 = wb.active
    ws1.title = "1 — Исправить названия"
    ws1.freeze_panes = "A3"
    ncols = 7
    add_title(ws1, f"Уровень 1 — Исправить названия в Dictionaries  ({len(l1)} записей)", ncols)
    headers = ["Id (GUID)", "Текущее название", "Правильное название",
               "Code", "IsDeleted", "IsCurrent", "Действие"]
    for col, h in enumerate(headers, 1):
        ws1.cell(row=2, column=col, value=h)
    style_header(ws1, 2, ncols, HDR_FILL_1)

    for r, row in enumerate(l1, 3):
        fill = WARN_FILL if not row["is_del"] else ERR_FILL
        write_cell(ws1, r, 1, row["guid"], fill)
        write_cell(ws1, r, 2, row["current"], fill)
        write_cell(ws1, r, 3, row["correct"], OK_FILL, bold=True)
        write_cell(ws1, r, 4, row["code"], fill)
        write_cell(ws1, r, 5, "Да" if row["is_del"] else "Нет", fill, center=True)
        write_cell(ws1, r, 6, "Да" if row["is_cur"] else "Нет", fill, center=True)
        write_cell(ws1, r, 7, "UPDATE Dictionaries.Name", fill, center=True)

    ws1.column_dimensions["A"].width = 38
    ws1.column_dimensions["B"].width = 42
    ws1.column_dimensions["C"].width = 52
    ws1.column_dimensions["D"].width = 10
    ws1.column_dimensions["E"].width = 10
    ws1.column_dimensions["F"].width = 10
    ws1.column_dimensions["G"].width = 28
    ws1.auto_filter.ref = f"A2:G{len(l1)+2}"

    # ── Лист 2: Уровень 2 — объединить GUID ──────────────────
    ws2 = wb.create_sheet("2 — Объединить GUID")
    ws2.freeze_panes = "A3"
    ncols2 = 13
    add_title(ws2, f"Уровень 2 — Объединить дублирующиеся GUID  ({len(l2)} пар)", ncols2)
    h2 = ["Страна (правильное)", "Удалить GUID", "Текущее название",
          "Оставить GUID", "Правильное название",
          "Cards\nCountry", "Cards\nCommission", "Cards\nLangSchool",
          "Placements", "Dogovors", "Movements", "Plans", "ИТОГО записей"]
    for col, h in enumerate(h2, 1):
        ws2.cell(row=2, column=col, value=h)
    style_header(ws2, 2, ncols2, HDR_FILL_2)

    for r, row in enumerate(l2, 3):
        c = row["counts"]
        total = row["total"]
        fill  = ERR_FILL if total > 0 else WARN_FILL
        write_cell(ws2, r,  1, row["correct_name"],  fill, bold=True)
        write_cell(ws2, r,  2, row["wrong_guid"],    ERR_FILL)
        write_cell(ws2, r,  3, row["wrong_name"],    ERR_FILL)
        write_cell(ws2, r,  4, row["correct_guid"],  OK_FILL)
        write_cell(ws2, r,  5, row["correct_name2"], OK_FILL)
        write_cell(ws2, r,  6, c.get("Cards.Country", 0),     fill, center=True)
        write_cell(ws2, r,  7, c.get("Cards.Commission", 0),  fill, center=True)
        write_cell(ws2, r,  8, c.get("Cards.LangSchool", 0),  fill, center=True)
        write_cell(ws2, r,  9, c.get("Placements", 0),        fill, center=True)
        write_cell(ws2, r, 10, c.get("Dogovors", 0),          fill, center=True)
        write_cell(ws2, r, 11, c.get("Movements", 0),         fill, center=True)
        write_cell(ws2, r, 12, c.get("Plans", 0),             fill, center=True)
        write_cell(ws2, r, 13, total, WARN_FILL if total > 0 else OK_FILL, bold=True, center=True)

    ws2.column_dimensions["A"].width = 42
    ws2.column_dimensions["B"].width = 38
    ws2.column_dimensions["C"].width = 38
    ws2.column_dimensions["D"].width = 38
    ws2.column_dimensions["E"].width = 38
    for col in "FGHIJKLM":
        ws2.column_dimensions[col].width = 12
    ws2.auto_filter.ref = f"A2:M{len(l2)+2}"

    # ── Лист 3: Уровень 3 — записи без Id ────────────────────
    ws3 = wb.create_sheet("3 — Записи без Id")
    ws3.freeze_panes = "A3"
    ncols3 = 6
    add_title(ws3, f"Уровень 3 — Записи без Id (NULL)  ({len(l3)} вариантов)", ncols3)
    h3 = ["Текущее значение", "Таблица", "Кол-во записей",
          "Правильное название", "GUID из справочника", "Статус"]
    for col, h in enumerate(h3, 1):
        ws3.cell(row=2, column=col, value=h)
    style_header(ws3, 2, ncols3, HDR_FILL_3)

    for r, row in enumerate(l3, 3):
        found  = row["status"] == "Найден GUID"
        fill_s = OK_FILL if found else ERR_FILL
        write_cell(ws3, r, 1, row["value"],   WARN_FILL)
        write_cell(ws3, r, 2, row["table"],   WARN_FILL)
        write_cell(ws3, r, 3, row["count"],   WARN_FILL, center=True)
        write_cell(ws3, r, 4, row["correct"], BLUE_FILL)
        write_cell(ws3, r, 5, row["guid"],    fill_s)
        write_cell(ws3, r, 6, row["status"],  fill_s, center=True)

    ws3.column_dimensions["A"].width = 50
    ws3.column_dimensions["B"].width = 22
    ws3.column_dimensions["C"].width = 14
    ws3.column_dimensions["D"].width = 50
    ws3.column_dimensions["E"].width = 38
    ws3.column_dimensions["F"].width = 18
    ws3.auto_filter.ref = f"A2:F{len(l3)+2}"

    # ── Лист 4: Сводка ────────────────────────────────────────
    ws4 = wb.create_sheet("Сводка")
    add_title(ws4, f"Сводка — план нормализации  ({datetime.now().strftime('%d.%m.%Y %H:%M')})", 3)

    summary = [
        ("", "", ""),
        ("Уровень", "Описание", "Кол-во"),
        ("1 — Исправить названия",
         "Неправильные названия в Dictionaries.Name",
         len(l1)),
        ("2 — Объединить GUID",
         "Одна страна — несколько GUID, нужно слить в один",
         len(l2)),
        ("3 — Привязать NULL записи",
         "Записи без DictionaryId — нужно найти и проставить GUID",
         sum(r["count"] for r in l3 if r["status"] == "Найден GUID")),
        ("3 — Не найдены в справочнике",
         "NULL записи у которых нет совпадения в Dictionaries",
         sum(r["count"] for r in l3 if r["status"] != "Найден GUID")),
        ("", "", ""),
        ("Порядок применения:", "", ""),
        ("Шаг 1", "Запустить full_normalize_countries.sql — блок АНАЛИЗ", ""),
        ("Шаг 2", "Проверить результаты анализа", ""),
        ("Шаг 3", "Раскомментировать BEGIN TRANSACTION", ""),
        ("Шаг 4", "Запустить — проверить PRINT-сообщения", ""),
        ("Шаг 5", "Выполнить COMMIT если всё верно", ""),
        ("Шаг 6", "Запустить export_countries_fast.py — проверить Excel", ""),
    ]

    ws4.column_dimensions["A"].width = 30
    ws4.column_dimensions["B"].width = 58
    ws4.column_dimensions["C"].width = 16

    for r, (a, b, c) in enumerate(summary, 2):
        ca = ws4.cell(row=r, column=1, value=a)
        cb = ws4.cell(row=r, column=2, value=b)
        cc = ws4.cell(row=r, column=3, value=c)
        for cell in (ca, cb, cc):
            cell.font   = NORM_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")
        if a == "Уровень":
            for cell in (ca, cb, cc):
                cell.font = HDR_FONT
                cell.fill = HDR_FILL_S
                cell.alignment = HDR_ALIGN
        elif a.startswith("1 —"):
            ca.fill = cb.fill = cc.fill = WARN_FILL
            cc.alignment = Alignment(horizontal="center", vertical="center")
        elif a.startswith("2 —"):
            ca.fill = cb.fill = cc.fill = ERR_FILL
            cc.alignment = Alignment(horizontal="center", vertical="center")
        elif a.startswith("3 —"):
            fill = OK_FILL if "Привязать" in a else ERR_FILL
            ca.fill = cb.fill = cc.fill = fill
            cc.alignment = Alignment(horizontal="center", vertical="center")
        elif a.startswith("Шаг"):
            ca.fill = BLUE_FILL
            ca.font = BOLD_FONT

    wb.save(output_path)
    print(f"\n✅ Excel сохранён: {output_path}")
    print(f"   Уровень 1 (названия):  {len(l1)} записей")
    print(f"   Уровень 2 (GUID дубли): {len(l2)} пар")
    print(f"   Уровень 3 (NULL записи): {len(l3)} вариантов")


# ============================================================
# MAIN
# ============================================================
def main():
    print("Подключаемся к БД (только чтение, Windows Auth)...")
    conn   = get_conn()
    cursor = conn.cursor()

    print("Собираем данные:")
    l1 = fetch_level1(cursor)
    l2 = fetch_level2(cursor)
    l3 = fetch_level3(cursor)

    conn.close()

    output = f"normalization_plan_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    print(f"\nФормируем Excel...")
    build_excel(l1, l2, l3, output)


if __name__ == "__main__":
    main()